import json
import csv
import io
from typing import List, Any
from app.modules.evaluation.committee.committee_models import (
    InvestmentCommitteeDecision, CommitteeReport
)
from app.modules.evaluation.committee.committee_queries import committee_dashboard


class CommitteeSerializer:
    """Handles JSON/CSV/Excel serialization, deserialization, and report generation for Investment Committee Decisions."""

    @staticmethod
    def to_json(obj: Any) -> str:
        """Serializes InvestmentCommitteeDecision or CommitteeReport to a JSON string."""
        if hasattr(obj, "model_dump_json"):
            return obj.model_dump_json()
        return json.dumps(obj)

    @staticmethod
    def from_json(json_str: str, target_class: Any = None) -> Any:
        """Deserializes a JSON string into InvestmentCommitteeDecision or CommitteeReport."""
        if target_class is not None:
            return target_class.model_validate_json(json_str)
        if '"findings"' in json_str and '"consensus"' in json_str:
            return CommitteeReport.model_validate_json(json_str)
        return InvestmentCommitteeDecision.model_validate_json(json_str)


    @staticmethod
    def to_csv(decisions: List[InvestmentCommitteeDecision]) -> str:
        """Exports committee decisions to a flat CSV format string."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "decision_id", "startup_id", "startup_name", "portfolio_rank",
            "recommendation", "decision_confidence", "investment_priority",
            "incubation_priority", "grant_priority", "pilot_priority", "review_window"
        ])
        for dec in decisions:
            writer.writerow([
                dec.decision_id,
                dec.startup_id,
                dec.startup_name,
                dec.portfolio_rank,
                dec.recommendation.value,
                dec.decision_confidence,
                dec.investment_priority.value,
                dec.incubation_priority.value,
                dec.grant_priority.value,
                dec.pilot_priority.value,
                dec.review_window
            ])
        return output.getvalue()

    @staticmethod
    def to_excel(decisions: List[InvestmentCommitteeDecision]) -> bytes:
        """Exports decisions and statistics to a styled multi-tab Excel spreadsheet (bytes)."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        wb = Workbook()
        ws_decisions = wb.active
        ws_decisions.title = "Committee Decisions"

        headers = [
            "Decision ID", "Startup ID", "Startup Name", "Portfolio Rank",
            "Recommendation", "Confidence", "Investment Priority",
            "Incubation Priority", "Grant Priority", "Pilot Priority", "Review Window"
        ]
        ws_decisions.append(headers)

        # Style header row
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(headers) + 1):
            cell = ws_decisions.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Add data
        for dec in decisions:
            ws_decisions.append([
                dec.decision_id,
                dec.startup_id,
                dec.startup_name,
                dec.portfolio_rank,
                dec.recommendation.value,
                dec.decision_confidence,
                dec.investment_priority.value,
                dec.incubation_priority.value,
                dec.grant_priority.value,
                dec.pilot_priority.value,
                dec.review_window
            ])

        # Style data rows
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        for row in range(2, len(decisions) + 2):
            for col in range(1, len(headers) + 1):
                cell = ws_decisions.cell(row=row, column=col)
                cell.border = thin_border
                if col in (4, 6):
                    cell.alignment = Alignment(horizontal="right")
                elif col in (1, 2, 5, 7, 8, 9, 10, 11):
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")

        # Stats Sheet
        ws_stats = wb.create_sheet(title="Dashboard Stats")
        ws_stats.append(["Metric", "Value"])
        ws_stats.cell(row=1, column=1).font = header_font
        ws_stats.cell(row=1, column=1).fill = header_fill
        ws_stats.cell(row=1, column=2).font = header_font
        ws_stats.cell(row=1, column=2).fill = header_fill

        from app.modules.evaluation.committee.committee_queries import statistics
        stats = statistics(decisions)
        
        ws_stats.append(["Total Decisions", stats["total_decisions"]])
        ws_stats.append(["Ready for Incubation", stats["ready_for_incubation"]])
        ws_stats.append(["Pending Due Diligence", stats["pending_due_diligence"]])
        ws_stats.append(["Average Confidence", stats["average_confidence"]])

        ws_stats.append([])
        ws_stats.append(["Recommendation Distribution", "Count"])
        row_offset = ws_stats.max_row
        ws_stats.cell(row=row_offset, column=1).font = header_font
        ws_stats.cell(row=row_offset, column=1).fill = header_fill
        ws_stats.cell(row=row_offset, column=2).font = header_font
        ws_stats.cell(row=row_offset, column=2).fill = header_fill
        for rec, count in stats["recommendation_distribution"].items():
            ws_stats.append([rec, count])

        ws_stats.append([])
        ws_stats.append(["Priority Distribution", "Count"])
        row_offset = ws_stats.max_row
        ws_stats.cell(row=row_offset, column=1).font = header_font
        ws_stats.cell(row=row_offset, column=1).fill = header_fill
        ws_stats.cell(row=row_offset, column=2).font = header_font
        ws_stats.cell(row=row_offset, column=2).fill = header_fill
        for prio, count in stats["priority_distribution"].items():
            ws_stats.append([prio, count])

        # Auto-fit columns
        for ws in (ws_decisions, ws_stats):
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def to_dashboard_json(decisions: List[InvestmentCommitteeDecision]) -> str:
        """Serializes dashboard views to JSON format."""
        return json.dumps(committee_dashboard(decisions))

    @staticmethod
    def to_markdown(obj: Any) -> str:
        """Generates a structured, styled Markdown committee decision report."""
        if isinstance(obj, InvestmentCommitteeDecision):
            # Legacy implementation
            md = []
            md.append(f"# Investment Committee Decision: {obj.decision_id}")
            md.append(f"**Startup Name:** {obj.startup_name} (ID: {obj.startup_id})")
            md.append(f"**Portfolio Rank:** {obj.portfolio_rank}")
            md.append(f"**Recommendation:** **{obj.recommendation.value}**")
            md.append(f"**Decision Confidence:** {obj.decision_confidence:.2f}")
            md.append(f"**Generated At:** {obj.generated_at}")
            md.append(f"**Review Window:** {obj.review_window}\n")
            
            md.append("## Executive Rationale")
            md.append(obj.decision_reasoning)
            md.append(f"\n*Notes: {obj.committee_notes}*\n")

            md.append("## Priorities")
            md.append("| Area | Priority Level |")
            md.append("| :--- | :--- |")
            md.append(f"| Investment | {obj.investment_priority.value} |")
            md.append(f"| Incubation | {obj.incubation_priority.value} |")
            md.append(f"| Grant | {obj.grant_priority.value} |")
            md.append(f"| Pilot/Sandbox | {obj.pilot_priority.value} |\n")

            md.append("## Blocking Risks")
            if obj.blocking_risks:
                for risk in obj.blocking_risks:
                    md.append(f"* {risk}")
            else:
                md.append("*No blocking risks identified.*")
            md.append("")

            md.append("## Required Due Diligence Checklist")
            md.append("| Category | Status | Blocking | Documents Required | Reason |")
            md.append("| :--- | :---: | :---: | :--- | :--- |")
            for item in obj.required_due_diligence:
                blocking_str = "Yes" if item.blocking else "No"
                docs_str = ", ".join(item.documents_required)
                md.append(f"| {item.category.value} | {item.status} | {blocking_str} | {docs_str} | {item.reason} |")
            md.append("")

            md.append("## Follow-up Questions")
            for q in obj.follow_up_questions:
                md.append(f"1. {q}")
            md.append("")

            md.append("## Lineage Metadata")
            md.append(f"* **Graph Hash:** `{obj.graph_hash}`")
            md.append(f"* **Portfolio Hash:** `{obj.portfolio_hash}`")
            md.append(f"* **Executive Assessment ID:** `{obj.executive_summary_id}`")
            md.append(f"* **Investment Assessment ID:** `{obj.investment_assessment_id}`")
            
            return "\n".join(md)

        # CommitteeReport implementation
        from app.modules.evaluation.committee.committee_templates import MARKDOWN_TEMPLATE
        report = obj
        
        findings_content = ""
        for f in report.findings:
            findings_content += f"### {f.title} *(Confidence: {f.confidence:.2f})*\n"
            findings_content += f"{f.description}\n\n"
            findings_content += f"*   **Domains:** {', '.join(f.domains)}\n"
            if f.supporting_observations:
                findings_content += f"*   **Observations:** {', '.join(f.supporting_observations)}\n"
            if f.supporting_claims:
                findings_content += f"*   **Claims:** {', '.join(f.supporting_claims)}\n"
            if f.supporting_evidence:
                findings_content += f"*   **Evidence:** {', '.join(f.supporting_evidence)}\n"
            findings_content += "\n"
            
        concerns_content = ""
        for c in report.concerns:
            concerns_content += f"### Concern: {c.concern_id} *(Severity: {c.severity}, Confidence: {c.confidence:.2f})*\n"
            concerns_content += f"{c.description}\n\n"
            concerns_content += f"*   **Affected Domains:** {', '.join(c.affected_domains)}\n"
            if c.linked_risks:
                concerns_content += f"*   **Linked Risks:** {', '.join(c.linked_risks)}\n"
            if c.linked_questions:
                concerns_content += f"*   **Linked Questions:** {', '.join(c.linked_questions)}\n"
            concerns_content += "\n"
            
        dependencies_content = ""
        for dep in report.consensus.cross_domain_dependencies:
            dependencies_content += f"*   **{dep.get('source_domain')}** -> **{dep.get('target_domain')}** ({dep.get('dependency_type')}): {dep.get('description')}\n"
        if not dependencies_content:
            dependencies_content = "*No cross-domain dependencies identified.*"
            
        followups_content = ""
        for q in report.recommendations.required_followups:
            followups_content += f"*   {q}\n"
        if not followups_content:
            followups_content = "*No follow-up questions required.*"
            
        documents_content = ""
        for doc in report.recommendations.required_documents:
            documents_content += f"*   {doc}\n"
        if not documents_content:
            documents_content = "*No additional documents required.*"
            
        return MARKDOWN_TEMPLATE.format(
            startup_name=report.decision.startup_name,
            decision_id=report.decision.decision_id,
            startup_id=report.decision.startup_id,
            workflow_reference=report.decision.workflow_reference,
            created_at=report.decision.created_at.isoformat() if hasattr(report.decision.created_at, "isoformat") else str(report.decision.created_at),
            overall_confidence=report.decision.overall_confidence,
            recommendation=report.recommendations.recommendation,
            priority=report.recommendations.priority,
            justification=report.recommendations.justification,
            overall_reasoning=report.decision.overall_reasoning,
            findings_content=findings_content.strip(),
            concerns_content=concerns_content.strip(),
            agreement_level=report.consensus.agreement_level,
            corroborated_count=len(report.consensus.corroborated_findings),
            contradictions_count=len(report.consensus.contradictions),
            unresolved_conflicts_count=len(report.consensus.unresolved_conflicts),
            dependencies_content=dependencies_content.strip(),
            followups_content=followups_content.strip(),
            documents_content=documents_content.strip()
        )

    @staticmethod
    def to_html(report: CommitteeReport) -> str:
        """Generates a structured, styled HTML committee report."""
        from app.modules.evaluation.committee.committee_templates import HTML_TEMPLATE
        
        findings_html = ""
        for f in report.findings:
            findings_html += f"<div style='margin-bottom:20px;'>"
            findings_html += f"<h3>{f.title} <span class='badge' style='background-color:#e8f0fe; color:#1a73e8;'>Confidence: {f.confidence:.2f}</span></h3>"
            findings_html += f"<p>{f.description}</p>"
            findings_html += f"<ul>"
            findings_html += f"<li><strong>Domains:</strong> {', '.join(f.domains)}</li>"
            if f.supporting_observations:
                findings_html += f"<li><strong>Observations:</strong> " + " ".join(f"<span class='trace-label'>{o}</span>" for o in f.supporting_observations) + "</li>"
            findings_html += f"</ul></div>"
            
        concerns_html = ""
        for c in report.concerns:
            concerns_html += f"<div style='margin-bottom:20px;'>"
            concerns_html += f"<h3>Concern {c.concern_id} <span class='badge' style='background-color:#fce8e6; color:#c5221f;'>Severity: {c.severity}</span></h3>"
            concerns_html += f"<p>{c.description}</p>"
            concerns_html += f"<ul>"
            concerns_html += f"<li><strong>Affected Domains:</strong> {', '.join(c.affected_domains)}</li>"
            if c.linked_risks:
                concerns_html += f"<li><strong>Linked Risks:</strong> " + " ".join(f"<span class='trace-label'>{r}</span>" for r in c.linked_risks) + "</li>"
            concerns_html += f"</ul></div>"
            
        documents_html = "".join(f"<li class='list-item'>{doc}</li>" for doc in report.recommendations.required_documents)
        if not documents_html:
            documents_html = "<li class='list-item'>No documents required.</li>"
            
        followups_html = "".join(f"<li class='list-item'>{q}</li>" for doc in report.recommendations.required_followups for q in [doc])
        if not followups_html:
            followups_html = "<li class='list-item'>No follow-up questions required.</li>"
            
        return HTML_TEMPLATE.format(
            startup_name=report.decision.startup_name,
            decision_id=report.decision.decision_id,
            startup_id=report.decision.startup_id,
            workflow_reference=report.decision.workflow_reference,
            overall_confidence=report.decision.overall_confidence,
            recommendation=report.recommendations.recommendation,
            priority=report.recommendations.priority,
            justification=report.recommendations.justification,
            overall_reasoning=report.decision.overall_reasoning,
            findings_html=findings_html,
            concerns_html=concerns_html,
            agreement_level=report.consensus.agreement_level,
            corroborated_count=len(report.consensus.corroborated_findings),
            contradictions_count=len(report.consensus.contradictions),
            unresolved_conflicts_count=len(report.consensus.unresolved_conflicts),
            documents_html=documents_html,
            followups_html=followups_html
        )

