import json
import csv
import io
from typing import Dict, Any
from app.modules.evaluation.portfolio.portfolio_models import Portfolio


class PortfolioSerializer:
    """Handles JSON serialization, deserialization, and exporting reports for TIE Startup Portfolios."""

    @staticmethod
    def to_json(portfolio: Portfolio) -> str:
        """Serializes the Portfolio to a JSON string."""
        return portfolio.model_dump_json(indent=2)

    @staticmethod
    def from_json(json_str: str) -> Portfolio:
        """Deserializes a Portfolio from a JSON string."""
        return Portfolio.model_validate_json(json_str)

    @staticmethod
    def to_csv(portfolio: Portfolio) -> str:
        """Exports rankings to a CSV format string."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "rank", "startup_id", "startup_name", "investment_score", "recommendation",
            "confidence", "percentile", "category", "ranking_reason", "graph_hash"
        ])
        for entry in portfolio.entries:
            writer.writerow([
                entry.rank,
                entry.startup_id,
                entry.startup_name,
                entry.investment_score,
                entry.recommendation,
                entry.confidence,
                entry.percentile,
                entry.category,
                entry.ranking_reason,
                entry.graph_hash
            ])
        return output.getvalue()

    @staticmethod
    def to_excel(portfolio: Portfolio) -> bytes:
        """Exports the report details to a styled Excel spreadsheet (bytes) using openpyxl."""
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        wb = Workbook()
        # Rankings Sheet
        ws_rankings = wb.active
        ws_rankings.title = "Rankings"

        # Headers
        headers = [
            "Rank", "Startup ID", "Startup Name", "Score", "Recommendation",
            "Confidence", "Percentile", "Category", "Ranking Reason", "Graph Hash"
        ]
        ws_rankings.append(headers)

        # Style header row
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_idx in range(1, len(headers) + 1):
            cell = ws_rankings.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Add data
        for entry in portfolio.entries:
            ws_rankings.append([
                entry.rank,
                entry.startup_id,
                entry.startup_name,
                entry.investment_score,
                entry.recommendation,
                entry.confidence,
                entry.percentile,
                entry.category,
                entry.ranking_reason,
                entry.graph_hash
            ])

        # Style data rows
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        for row in range(2, len(portfolio.entries) + 2):
            for col in range(1, len(headers) + 1):
                cell = ws_rankings.cell(row=row, column=col)
                cell.border = thin_border
                if col in (1, 4, 6, 7):
                    cell.alignment = Alignment(horizontal="right")
                elif col in (2, 5, 10):
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")

        # Stats Sheet
        ws_stats = wb.create_sheet(title="Statistics")
        ws_stats.append(["Metric", "Value"])
        ws_stats.cell(row=1, column=1).font = header_font
        ws_stats.cell(row=1, column=1).fill = header_fill
        ws_stats.cell(row=1, column=2).font = header_font
        ws_stats.cell(row=1, column=2).fill = header_fill

        stats = portfolio.statistics
        stats_data = [
            ("Portfolio Size", stats.portfolio_size),
            ("Average Score", stats.average_score),
            ("Median Score", stats.median_score),
            ("Highest Score", stats.highest_score),
            ("Lowest Score", stats.lowest_score),
            ("Average Confidence", stats.average_confidence),
        ]
        for metric, val in stats_data:
            ws_stats.append([metric, val])

        # Add distributions in stats sheet
        ws_stats.append([])
        ws_stats.append(["Recommendation", "Count"])
        row_offset = ws_stats.max_row
        ws_stats.cell(row=row_offset, column=1).font = header_font
        ws_stats.cell(row=row_offset, column=1).fill = header_fill
        ws_stats.cell(row=row_offset, column=2).font = header_font
        ws_stats.cell(row=row_offset, column=2).fill = header_fill
        for rec, count in stats.recommendation_distribution.items():
            ws_stats.append([rec, count])

        ws_stats.append([])
        ws_stats.append(["Category", "Count"])
        row_offset = ws_stats.max_row
        ws_stats.cell(row=row_offset, column=1).font = header_font
        ws_stats.cell(row=row_offset, column=1).fill = header_fill
        ws_stats.cell(row=row_offset, column=2).font = header_font
        ws_stats.cell(row=row_offset, column=2).fill = header_fill
        for cat, count in stats.category_distribution.items():
            ws_stats.append([cat, count])

        # Column width auto-fit
        for ws in (ws_rankings, ws_stats):
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                col_letter = col[0].column_letter
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def to_dashboard_json(portfolio: Portfolio) -> str:
        """Exports summary metrics and top performers suitable for frontend dashboard widgets."""
        stats = portfolio.statistics
        return json.dumps({
            "portfolio_id": portfolio.portfolio_id,
            "generated_at": portfolio.generated_at,
            "size": stats.portfolio_size,
            "average_score": stats.average_score,
            "median_score": stats.median_score,
            "highest_score": stats.highest_score,
            "lowest_score": stats.lowest_score,
            "average_confidence": stats.average_confidence,
            "recommendation_distribution": stats.recommendation_distribution,
            "category_distribution": stats.category_distribution,
            "top_performers": [
                {
                    "rank": entry.rank,
                    "startup_id": entry.startup_id,
                    "startup_name": entry.startup_name,
                    "score": entry.investment_score,
                    "recommendation": entry.recommendation,
                    "percentile": entry.percentile,
                    "ranking_reason": entry.ranking_reason
                }
                for entry in portfolio.entries[:5]
            ]
        }, indent=2)

    @staticmethod
    def to_markdown_summary(portfolio: Portfolio) -> str:
        """Generates a beautiful Markdown summary report of the portfolio rankings."""
        stats = portfolio.statistics
        md = []
        md.append(f"# Portfolio Intelligence Report: {portfolio.portfolio_id}")
        md.append(f"**Generated At:** {portfolio.generated_at}")
        md.append(f"**Engine Version:** {portfolio.engine_version}\n")
        
        md.append("## Executive Statistics")
        md.append("| Metric | Value |")
        md.append("| :--- | :--- |")
        md.append(f"| Portfolio Size | {stats.portfolio_size} startups |")
        md.append(f"| Average Score | {stats.average_score:.2f} |")
        md.append(f"| Median Score | {stats.median_score:.2f} |")
        md.append(f"| Highest Score | {stats.highest_score:.2f} |")
        md.append(f"| Lowest Score | {stats.lowest_score:.2f} |")
        md.append(f"| Average Confidence | {stats.average_confidence:.4f} |\n")
        
        md.append("## Startup Rankings")
        md.append("| Rank | Startup ID | Startup Name | Score | Recommendation | Confidence | Percentile | Category | Ranking Reason |")
        md.append("| ---: | :--- | :--- | ---: | :--- | ---: | ---: | :--- | :--- |")
        for entry in portfolio.entries:
            md.append(
                f"| {entry.rank} | {entry.startup_id} | {entry.startup_name} | "
                f"{entry.investment_score:.2f} | {entry.recommendation} | {entry.confidence:.4f} | "
                f"{entry.percentile:.2f}% | {entry.category} | {entry.ranking_reason} |"
            )
            
        md.append("\n## Recommendation Distribution")
        for rec, count in stats.recommendation_distribution.items():
            md.append(f"* **{rec}**: {count}")
            
        md.append("\n## Category Distribution")
        for cat, count in stats.category_distribution.items():
            md.append(f"* **{cat}**: {count}")
            
        return "\n".join(md)
