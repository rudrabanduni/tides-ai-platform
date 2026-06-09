from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any
from uuid import UUID, uuid4
from zipfile import BadZipFile

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from app.core.enums import StartupStatus
from app.core.exceptions import ValidationError
from app.db.mixins import utcnow
from app.modules.audit.service import AuditService
from app.modules.founders.models import Founder
from app.modules.intake.schemas import IntakeRowError, IntakeUploadSummary
from app.modules.startup_profiles.models import StartupProfile, StartupProfileVersion
from app.modules.startups.models import StartupApplication, StartupStatusHistory


EXPECTED_COLUMNS: dict[str, str] = {
    "Startup Name": "startup_name",
    "Founder Name": "founder_name",
    "Founder Email": "founder_email",
    "Problem Statement": "problem_statement",
    "Solution": "solution_summary",
    "Target Market": "target_market",
    "Startup Stage": "stage",
}

REQUIRED_FIELDS: dict[str, str] = {
    "startup_name": "Startup Name",
    "founder_name": "Founder Name",
    "founder_email": "Founder Email",
}


@dataclass(frozen=True)
class IntakeRow:
    row_number: int
    startup_name: str | None
    founder_name: str | None
    founder_email: str | None
    problem_statement: str | None
    solution_summary: str | None
    target_market: str | None
    stage: str | None


class ExcelIntakeService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self.audit = AuditService(db)

    def import_applications_from_xlsx(
        self,
        *,
        filename: str,
        content: bytes,
        actor_id: UUID | None,
    ) -> IntakeUploadSummary:
        import_id = str(uuid4())
        workbook = None
        try:
            self._validate_file(filename, content)
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
            worksheet = workbook.active
            header_map = self._build_header_map(worksheet)
            summary = self._process_rows(
                worksheet=worksheet,
                header_map=header_map,
                filename=filename,
                import_id=import_id,
                actor_id=actor_id,
            )
            self.audit.log(
                actor_id=actor_id,
                entity_type="excel_import",
                entity_id=import_id,
                action="excel_import_completed",
                details={
                    "filename": filename,
                    "total_rows": summary.total_rows,
                    "successful_rows": summary.successful_rows,
                    "failed_rows": summary.failed_rows,
                    "errors": [error.model_dump() for error in summary.errors],
                },
            )
            self.db.commit()
            return summary
        except (BadZipFile, InvalidFileException) as exc:
            self._log_failed_import(
                actor_id=actor_id,
                import_id=import_id,
                filename=filename,
                error="Uploaded file is not a valid .xlsx workbook",
            )
            raise ValidationError("Uploaded file is not a valid .xlsx workbook") from exc
        except Exception as exc:
            self._log_failed_import(
                actor_id=actor_id,
                import_id=import_id,
                filename=filename,
                error=str(getattr(exc, "detail", None) or exc),
            )
            raise
        finally:
            if workbook is not None:
                workbook.close()

    def _log_failed_import(self, *, actor_id: UUID | None, import_id: str, filename: str, error: str) -> None:
        self.db.rollback()
        self.audit.log(
            actor_id=actor_id,
            entity_type="excel_import",
            entity_id=import_id,
            action="excel_import_failed",
            details={"filename": filename, "error": error},
        )
        self.db.commit()

    def _process_rows(
        self,
        *,
        worksheet: Worksheet,
        header_map: dict[str, int],
        filename: str,
        import_id: str,
        actor_id: UUID | None,
    ) -> IntakeUploadSummary:
        total_rows = 0
        successful_rows = 0
        errors: list[IntakeRowError] = []
        seen_founder_emails: set[str] = set()

        for row_number, raw_values in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            if self._is_empty_row(raw_values):
                continue

            total_rows += 1
            intake_row = self._row_from_values(row_number, raw_values, header_map)
            row_errors = self._validate_row(intake_row, seen_founder_emails)
            if row_errors:
                errors.extend(row_errors)
                continue

            founder_email = intake_row.founder_email.lower() if intake_row.founder_email else ""
            seen_founder_emails.add(founder_email)

            try:
                self._create_records_for_row(
                    intake_row,
                    filename=filename,
                    import_id=import_id,
                    actor_id=actor_id,
                )
            except SQLAlchemyError as exc:
                self.db.rollback()
                errors.append(
                    IntakeRowError(
                        row_number=row_number,
                        startup_name=intake_row.startup_name,
                        message=f"Database error while importing row: {exc.__class__.__name__}",
                    )
                )
                continue

            successful_rows += 1

        return IntakeUploadSummary(
            total_rows=total_rows,
            successful_rows=successful_rows,
            failed_rows=len(errors_by_row(errors)),
            errors=errors,
        )

    def _create_records_for_row(
        self,
        intake_row: IntakeRow,
        *,
        filename: str,
        import_id: str,
        actor_id: UUID | None,
    ) -> None:
        startup = StartupApplication(
            startup_name=intake_row.startup_name or "",
            stage=intake_row.stage,
            current_status=StartupStatus.SUBMITTED,
            problem_statement=intake_row.problem_statement,
            solution_summary=intake_row.solution_summary,
            target_market=intake_row.target_market,
            created_by=actor_id,
            submitted_at=utcnow(),
        )
        self.db.add(startup)
        self.db.flush()

        self.db.add(
            StartupStatusHistory(
                startup_id=startup.id,
                old_status=None,
                new_status=StartupStatus.SUBMITTED,
                changed_by=actor_id,
                reason=f"Imported from Excel file {filename}, row {intake_row.row_number}",
            )
        )

        founder = Founder(
            startup_id=startup.id,
            name=intake_row.founder_name or "",
            email=intake_row.founder_email,
        )
        self.db.add(founder)

        profile = StartupProfile(
            startup_id=startup.id,
            problem_statement=intake_row.problem_statement,
            solution_summary=intake_row.solution_summary,
            target_market=intake_row.target_market,
        )
        self.db.add(profile)
        self.db.flush()

        self.db.add(
            StartupProfileVersion(
                startup_profile_id=profile.id,
                version_number=1,
                profile_snapshot=self._profile_snapshot(profile, startup_id=startup.id),
            )
        )

        self.audit.log(
            actor_id=actor_id,
            entity_type="startup",
            entity_id=startup.id,
            action="startup_imported_from_excel",
            details={
                "import_id": import_id,
                "filename": filename,
                "row_number": intake_row.row_number,
                "startup_name": startup.startup_name,
                "founder_email": founder.email,
            },
        )
        self.audit.log(
            actor_id=actor_id,
            entity_type="excel_import_row",
            entity_id=f"{import_id}:{intake_row.row_number}",
            action="excel_import_row_succeeded",
            details={
                "startup_id": str(startup.id),
                "filename": filename,
                "row_number": intake_row.row_number,
            },
        )
        self.db.commit()

    def _validate_row(self, intake_row: IntakeRow, seen_founder_emails: set[str]) -> list[IntakeRowError]:
        errors: list[IntakeRowError] = []
        row_data = intake_row.__dict__

        for field_name, column_name in REQUIRED_FIELDS.items():
            if not row_data.get(field_name):
                errors.append(
                    IntakeRowError(
                        row_number=intake_row.row_number,
                        startup_name=intake_row.startup_name,
                        field=column_name,
                        message=f"{column_name} is required",
                    )
                )

        normalized_email = self._normalize_email(intake_row.founder_email)
        if intake_row.founder_email and normalized_email is None:
            errors.append(
                IntakeRowError(
                    row_number=intake_row.row_number,
                    startup_name=intake_row.startup_name,
                    field="Founder Email",
                    message="Founder Email is invalid",
                )
            )
        elif normalized_email:
            if normalized_email in seen_founder_emails:
                errors.append(
                    IntakeRowError(
                        row_number=intake_row.row_number,
                        startup_name=intake_row.startup_name,
                        field="Founder Email",
                        message="Founder Email is duplicated within the uploaded file",
                    )
                )
            elif self._founder_email_exists(normalized_email):
                errors.append(
                    IntakeRowError(
                        row_number=intake_row.row_number,
                        startup_name=intake_row.startup_name,
                        field="Founder Email",
                        message="Founder Email already exists in the platform",
                    )
                )

        return errors

    def _founder_email_exists(self, email: str) -> bool:
        statement = select(Founder.id).where(func.lower(Founder.email) == email.lower())
        return self.db.scalar(statement) is not None

    @staticmethod
    def _normalize_email(email: str | None) -> str | None:
        if not email:
            return None
        try:
            result = validate_email(email, check_deliverability=False)
        except EmailNotValidError:
            return None
        return result.normalized.lower()

    @staticmethod
    def _row_from_values(row_number: int, raw_values: tuple[Any, ...], header_map: dict[str, int]) -> IntakeRow:
        values = {
            field_name: normalize_cell(raw_values[column_index] if column_index < len(raw_values) else None)
            for field_name, column_index in header_map.items()
        }
        founder_email = values.get("founder_email")
        if founder_email:
            founder_email = founder_email.lower()
        return IntakeRow(
            row_number=row_number,
            startup_name=values.get("startup_name"),
            founder_name=values.get("founder_name"),
            founder_email=founder_email,
            problem_statement=values.get("problem_statement"),
            solution_summary=values.get("solution_summary"),
            target_market=values.get("target_market"),
            stage=values.get("stage"),
        )

    @staticmethod
    def _build_header_map(worksheet: Worksheet) -> dict[str, int]:
        first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first_row:
            raise ValidationError("Excel file must include a header row")

        normalized_headers = {
            normalize_header(header): index
            for index, header in enumerate(first_row)
            if normalize_header(header)
        }

        missing_columns = [
            column_name
            for column_name in EXPECTED_COLUMNS
            if normalize_header(column_name) not in normalized_headers
        ]
        if missing_columns:
            raise ValidationError(f"Excel file is missing required columns: {', '.join(missing_columns)}")

        return {
            field_name: normalized_headers[normalize_header(column_name)]
            for column_name, field_name in EXPECTED_COLUMNS.items()
        }

    @staticmethod
    def _validate_file(filename: str, content: bytes) -> None:
        if not filename.lower().endswith(".xlsx"):
            raise ValidationError("Only .xlsx files are supported")
        if not content:
            raise ValidationError("Uploaded Excel file is empty")

    @staticmethod
    def _is_empty_row(raw_values: tuple[Any, ...]) -> bool:
        return all(normalize_cell(value) is None for value in raw_values)

    @staticmethod
    def _profile_snapshot(profile: StartupProfile, *, startup_id: UUID) -> dict[str, str | None]:
        return {
            "startup_id": str(startup_id),
            "problem_statement": profile.problem_statement,
            "solution_summary": profile.solution_summary,
            "target_market": profile.target_market,
            "business_model": profile.business_model,
            "technology_summary": profile.technology_summary,
            "traction_summary": profile.traction_summary,
            "funding_summary": profile.funding_summary,
            "ip_summary": profile.ip_summary,
        }


def normalize_header(value: Any) -> str:
    return " ".join(str(value or "").strip().lower().split())


def normalize_cell(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def errors_by_row(errors: list[IntakeRowError]) -> set[int]:
    return {error.row_number for error in errors}
