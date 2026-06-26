import os
import tempfile
import pytest
from uuid import uuid4
from sqlalchemy.orm import Session

from app.db.base import Base
from app.modules.documents.models import Document
from app.core.enums import DocumentProcessingStatus, DocumentType
from app.modules.startups.models import StartupApplication
from app.modules.intelligence.classification import classify_document, DocumentClassification
from app.modules.intelligence.parsers import registry, ParserManager, TxtParser, PdfParser, DocxParser, PptxParser, XlsxParser
from app.modules.intelligence.events import dispatcher, Event
from app.modules.intelligence.service import IntelligenceService
from app.modules.intelligence.models import PipelineStatus

# Setup mock imports for test generations
from reportlab.pdfgen import canvas
import docx
import pptx
import openpyxl


@pytest.fixture()
def db_session() -> Session:
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    from sqlalchemy.pool import StaticPool

    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(engine)
    session = TestingSessionLocal()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(engine)
        engine.dispose()


def test_document_classification() -> None:
    # 1. Pitch Decks
    assert classify_document("startup_pitch_deck.pdf") == DocumentClassification.PITCH_DECK
    assert classify_document("presentation.pptx") == DocumentClassification.PITCH_DECK
    assert classify_document("demo_slides.ppt") == DocumentClassification.PITCH_DECK

    # 2. Business Plans
    assert classify_document("business_proposal.docx") == DocumentClassification.BUSINESS_PLAN
    assert classify_document("strategic_plan.pdf") == DocumentClassification.BUSINESS_PLAN

    # 3. Financial Statements
    assert classify_document("balance_sheet.xlsx") == DocumentClassification.FINANCIAL_STATEMENT
    assert classify_document("financials.csv") == DocumentClassification.FINANCIAL_STATEMENT
    assert classify_document("q4_revenue_report.pdf") == DocumentClassification.FINANCIAL_STATEMENT

    # 4. Patents & Research Papers
    assert classify_document("invention_patent_claims.pdf") == DocumentClassification.PATENT
    assert classify_document("machine_learning_paper.pdf") == DocumentClassification.RESEARCH_PAPER

    # 5. Excel Intake
    assert classify_document("intake_form.xlsx") == DocumentClassification.EXCEL_INTAKE

    # 6. Fallback
    assert classify_document("random_file.xyz") == DocumentClassification.UNKNOWN


def test_parser_registry() -> None:
    # Verify that all standard parsers are registered
    assert registry.get_parser("text/plain", ".txt") is not None
    assert registry.get_parser("application/pdf", ".pdf") is not None
    assert registry.get_parser("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx") is not None
    assert registry.get_parser("application/vnd.openxmlformats-officedocument.presentationml.presentation", ".pptx") is not None
    assert registry.get_parser("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx") is not None

    # Unknown ext should return None
    assert registry.get_parser("unknown/mime", ".unknown") is None


def test_txt_parser() -> None:
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as tmp:
        tmp.write("This is a simple text file for testing purposes.")
        tmp_path = tmp.name

    try:
        parser = TxtParser()
        assert parser.supports("text/plain", ".txt") is True

        res = parser.parse(tmp_path)
        assert res["normalized_text"] == "This is a simple text file for testing purposes."
        assert res["page_count"] == 1
        assert res["word_count"] == 9
        assert res["language"] == "en"
        assert res["metadata"]["title"] == os.path.basename(tmp_path)
    finally:
        os.unlink(tmp_path)


def test_pdf_parser() -> None:
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        # Generate simple PDF
        c = canvas.Canvas(tmp_path)
        c.drawString(100, 750, "Hello, this is a test PDF.")
        c.save()

        parser = PdfParser()
        assert parser.supports("application/pdf", ".pdf") is True

        res = parser.parse(tmp_path)
        assert "Hello, this is a test PDF." in res["normalized_text"]
        assert res["page_count"] == 1
        assert res["word_count"] > 0
        assert res["language"] == "en"
        assert res["metadata"]["author"] is not None
    finally:
        os.unlink(tmp_path)


def test_docx_parser() -> None:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        # Generate simple DOCX
        doc = docx.Document()
        doc.add_paragraph("This is DOCX document content.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header1"
        table.cell(0, 1).text = "Header2"
        table.cell(1, 0).text = "Value1"
        table.cell(1, 1).text = "Value2"
        doc.save(tmp_path)

        parser = DocxParser()
        assert parser.supports("application/vnd.openxmlformats-officedocument.wordprocessingml.document", ".docx") is True

        res = parser.parse(tmp_path)
        assert "This is DOCX document content." in res["normalized_text"]
        assert "Header1 | Header2" in res["normalized_text"]
        assert len(res["tables"]) == 1
        assert res["tables"][0][0] == ["Header1", "Header2"]
        assert res["page_count"] >= 1
    finally:
        os.unlink(tmp_path)


def test_pptx_parser() -> None:
    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        # Generate simple PPTX
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = "TIE Presentation Title"
        prs.save(tmp_path)

        parser = PptxParser()
        assert parser.supports("application/vnd.openxmlformats-officedocument.presentationml.presentation", ".pptx") is True

        res = parser.parse(tmp_path)
        assert "TIE Presentation Title" in res["normalized_text"]
        assert res["page_count"] == 1
    finally:
        os.unlink(tmp_path)


def test_xlsx_parser() -> None:
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        # Generate simple XLSX
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SummarySheet"
        ws["A1"] = "Intake Key"
        ws["B1"] = "Value"
        ws["A2"] = "trl"
        ws["B2"] = 4
        wb.save(tmp_path)

        parser = XlsxParser()
        assert parser.supports("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", ".xlsx") is True

        res = parser.parse(tmp_path)
        assert "Intake Key | Value" in res["normalized_text"]
        assert "trl | 4" in res["normalized_text"]
        assert res["page_count"] == 1
        assert len(res["tables"]) == 1
        assert res["tables"][0]["sheet_name"] == "SummarySheet"
    finally:
        os.unlink(tmp_path)


def test_parser_manager_unsupported_and_corrupted() -> None:
    # Unsupported file extension (file must exist to bypass FileNotFoundError check)
    with tempfile.NamedTemporaryFile(suffix=".abc", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        with pytest.raises(ValueError, match="No parser registered"):
            ParserManager.parse_file(tmp_path, "unknown/mime")
    finally:
        os.unlink(tmp_path)

    # Non-existent file
    with pytest.raises(FileNotFoundError):
        ParserManager.parse_file("missing_file_path.pdf")

    # Corrupted PDF
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(b"corrupted pdf bytes")
        tmp_path = tmp.name

    try:
        with pytest.raises(ValueError, match="Parsing failed"):
            ParserManager.parse_file(tmp_path, "application/pdf")
    finally:
        os.unlink(tmp_path)


def test_document_pipeline_integration(db_session: Session) -> None:
    # Setup startup database record
    startup = StartupApplication(
        startup_name="TIE Integration Startup",
        stage="MVP",
        current_status="Submitted"
    )
    db_session.add(startup)
    db_session.flush()

    # Create dummy text document file
    with tempfile.NamedTemporaryFile(suffix=".txt", delete=False, mode="w", encoding="utf-8") as tmp:
        tmp.write("Pitch Deck text: We are building next-gen AI systems.")
        tmp_path = tmp.name

    # Create document database record
    doc = Document(
        startup_id=startup.id,
        document_type=DocumentType.OTHER,
        original_filename="startup_pitch_deck.txt",
        stored_filename=os.path.basename(tmp_path),
        file_path=tmp_path,
        content_type="text/plain",
        file_size=len("Pitch Deck text: We are building next-gen AI systems."),
        processing_status=DocumentProcessingStatus.UPLOADED
    )
    db_session.add(doc)
    db_session.flush()

    # Event tracking
    events_dispatched = []
    def tracker(event: Event) -> None:
        events_dispatched.append(event)

    dispatcher.register("DocumentRegistered", tracker)
    dispatcher.register("DocumentClassified", tracker)
    dispatcher.register("DocumentParsed", tracker)
    dispatcher.register("DocumentFailed", tracker)

    service = IntelligenceService(db_session)
    
    try:
        # Run processing
        meta = service.process_document(doc.id)

        # Check DB updates
        db_session.refresh(doc)
        assert doc.processing_status == DocumentProcessingStatus.PARSED
        assert "next-gen AI" in doc.parsed_text
        assert doc.document_type == DocumentType.PITCH_DECK  # Classified as Pitch Deck due to filename

        # Verify sidecar meta file creation
        meta_path = tmp_path + ".meta.json"
        assert os.path.exists(meta_path)
        loaded_meta = service.get_document_metadata(tmp_path)
        assert loaded_meta is not None
        assert loaded_meta["classification"] == "Pitch Deck"
        assert loaded_meta["word_count"] > 0
        assert loaded_meta["page_count"] == 1

        # Check pipeline status in database
        status_record = service.get_processing_status(startup.id, "document_intelligence")
        assert status_record is not None
        assert status_record.status == PipelineStatus.COMPLETED
        assert status_record.current_stage == "PARSED"

        # Check event emissions
        event_names = [e.event_name for e in events_dispatched]
        assert "DocumentRegistered" in event_names
        assert "DocumentClassified" in event_names
        assert "DocumentParsed" in event_names
    finally:
        # Cleanup sidecar and file
        os.unlink(tmp_path)
        os.unlink(tmp_path + ".meta.json")


def test_document_pipeline_failure_integration(db_session: Session) -> None:
    # Setup startup database record
    startup = StartupApplication(
        startup_name="TIE Failure Startup",
        stage="MVP",
        current_status="Submitted"
    )
    db_session.add(startup)
    db_session.flush()

    # Create dummy corrupted PDF document file
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
        tmp.write(b"invalid pdf data")
        tmp_path = tmp.name

    # Create document database record
    doc = Document(
        startup_id=startup.id,
        document_type=DocumentType.OTHER,
        original_filename="corrupted_doc.pdf",
        stored_filename=os.path.basename(tmp_path),
        file_path=tmp_path,
        content_type="application/pdf",
        file_size=16,
        processing_status=DocumentProcessingStatus.UPLOADED
    )
    db_session.add(doc)
    db_session.flush()

    # Event tracking
    events_dispatched = []
    def tracker(event: Event) -> None:
        events_dispatched.append(event)

    dispatcher.register("DocumentFailed", tracker)

    service = IntelligenceService(db_session)
    
    try:
        # Run processing and expect ValueError
        with pytest.raises(ValueError):
            service.process_document(doc.id)

        # Check DB updates
        db_session.refresh(doc)
        assert doc.processing_status == DocumentProcessingStatus.FAILED

        # Check pipeline status in database is FAILED
        status_record = service.get_processing_status(startup.id, "document_intelligence")
        assert status_record is not None
        assert status_record.status == PipelineStatus.FAILED
        assert status_record.current_stage == "FAILED"
        assert "Parsing failed" in status_record.last_error

        # Check event emissions
        event_names = [e.event_name for e in events_dispatched]
        assert "DocumentFailed" in event_names
    finally:
        os.unlink(tmp_path)
